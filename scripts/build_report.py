"""Build the DOCX report delivered for Laboratory 06."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
CAPTURES = ROOT / "docs" / "capturas"
OUTPUT = ROOT / "docs" / "ENTREGABLE_LABORATORIO_06.docx"

PRIMARY = RGBColor(0x0B, 0x3D, 0x5C)
ACCENT = RGBColor(0xD9, 0x89, 0x1C)
MUTED = RGBColor(0x66, 0x77, 0x86)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F5F7")
    p_pr.append(shd)
    run = paragraph.add_run(text.strip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(8.2)
    run.font.color.rgb = RGBColor(0x23, 0x31, 0x3D)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def add_figure(document, filename, caption, width=6.2):
    document.add_picture(str(CAPTURES / filename), width=Inches(width))
    last = document.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(document, caption)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, "0B3D5C")
        if widths:
            cell.width = Cm(widths[index])
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(value)
            run.font.size = Pt(8.5)
            if widths:
                cells[index].width = Cm(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def style_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 28, PRIMARY),
        ("Heading 1", 16, PRIMARY),
        ("Heading 2", 12.5, ACCENT),
        ("Heading 3", 11, PRIMARY),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(
        "Laboratorio 06 - Portal de Noticias con Django | Página "
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED
    add_page_number(footer)


def cover(document):
    for _ in range(3):
        document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("TECSUP - DEPARTAMENTO DE TECNOLOGÍA DIGITAL")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("INFORME DE LABORATORIO N° 06")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = PRIMARY

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(
        "Motor de Plantillas con Django:\nHerencia, Componentes y Gestión de Contenido"
    )
    run.font.size = Pt(15)
    run.font.color.rgb = ACCENT
    run.bold = True

    document.add_paragraph()
    add_table(
        document,
        ["Dato", "Valor"],
        [
            ("Estudiante", "Galindo Huaman Edu Edward"),
            ("Curso", "Desarrollo de Aplicaciones Empresariales"),
            ("Sección / Grupo", "4-C24-A"),
            ("Fecha de ejecución", "Semana 06 - Ciclo 2026-02"),
            ("Proyecto", "news_portal"),
            ("Tecnologías", "Django 5.2.17, Pillow 12.3.0, SQLite"),
        ],
        widths=[5, 10],
    )
    document.add_page_break()


def build():
    document = Document()
    style_document(document)
    cover(document)

    document.add_heading("1. Capacidades y normas de desarrollo", level=1)
    document.add_paragraph(
        "En esta sesión se implementó una arquitectura web modular con Django "
        "empleando herencia de plantillas, fragmentos reutilizables, "
        "renderizado dinámico de datos, filtros de formato y gestión de "
        "contenido desde el panel de administración."
    )
    for item in (
        "Motor de plantillas: herencia con {% extends %} y componentes con "
        "{% include %}.",
        "Renderizado dinámico: variables, {% for %}, {% if %}, {% empty %}, "
        "date y truncatewords.",
        "Gestión de contenido: categorías, autores y noticias administradas "
        "desde Django Admin.",
        "Seguridad: verificación del auto-escapado de HTML contra XSS.",
        "Estándares: código Python en inglés, PEP 8 y documentación en español.",
    ):
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("2. Estructura del proyecto", level=1)
    document.add_paragraph(
        "El proyecto news_portal separa configuración, aplicación, plantillas, "
        "estáticos, medios, documentación y pruebas automatizadas."
    )
    add_code(
        document,
        """
news_portal/
├── config/
│   ├── settings.py
│   └── urls.py
├── news/
│   ├── management/commands/seed_demo.py
│   ├── migrations/0001_initial.py
│   ├── admin.py
│   ├── models.py
│   ├── tests.py
│   ├── urls.py
│   └── views.py
├── media/{articles,authors}/
├── static/css/styles.css
├── templates/
│   ├── base.html
│   └── news/
│       ├── _article_card.html
│       ├── article_detail.html
│       ├── category_list.html
│       └── home.html
├── docs/capturas/
├── manage.py
└── requirements.txt
""",
    )
    add_figure(
        document,
        "01_portada.png",
        "Figura 1. Portada del portal con seis noticias y barra lateral.",
    )

    document.add_heading("3. Desarrollo del procedimiento paso a paso", level=1)

    document.add_heading(
        "Pasos 1 y 2: proyecto, aplicación y archivos estáticos/medios",
        level=2,
    )
    document.add_paragraph(
        "Se creó el proyecto con django-admin startproject config . y la "
        "aplicación news. Se instaló Pillow y se registró news.apps.NewsConfig "
        "en INSTALLED_APPS. También se configuraron las plantillas, los "
        "estáticos y los medios."
    )
    add_code(
        document,
        """
INSTALLED_APPS = [
    'django.contrib.admin',
    'django.contrib.auth',
    'django.contrib.contenttypes',
    'django.contrib.sessions',
    'django.contrib.messages',
    'django.contrib.staticfiles',
    'news.apps.NewsConfig',
]

TEMPLATES[0]['DIRS'] = [BASE_DIR / 'templates']
STATIC_URL = 'static/'
STATICFILES_DIRS = [BASE_DIR / 'static']
MEDIA_URL = '/media/'
MEDIA_ROOT = BASE_DIR / 'media'

# config/urls.py
if settings.DEBUG:
    urlpatterns += static(
        settings.MEDIA_URL,
        document_root=settings.MEDIA_ROOT,
    )
""",
    )

    document.add_heading(
        "Paso 3: modelos, migraciones y superusuario",
        level=2,
    )
    document.add_paragraph(
        "Se declararon Category, Author y Article con imagen destacada, fecha "
        "de publicación y relaciones FK/M2M. Luego se ejecutaron makemigrations, "
        "migrate y createsuperuser."
    )
    add_code(
        document,
        """
class Category(models.Model):
    name = models.CharField(max_length=100, unique=True)
    slug = models.SlugField(max_length=100, unique=True)

class Author(models.Model):
    first_name = models.CharField(max_length=100)
    last_name = models.CharField(max_length=100)
    email = models.EmailField(unique=True)
    bio = models.TextField(blank=True)
    profile_picture = models.ImageField(
        upload_to='authors/', blank=True, null=True
    )

class Article(models.Model):
    title = models.CharField(max_length=200)
    slug = models.SlugField(max_length=200, unique=True)
    summary = models.TextField(help_text='Brief description for news cards')
    content = models.TextField()
    featured_image = models.ImageField(upload_to='articles/')
    published_date = models.DateTimeField(auto_now_add=True)
    is_published = models.BooleanField(default=True)
    author = models.ForeignKey(
        Author, on_delete=models.CASCADE, related_name='articles'
    )
    categories = models.ManyToManyField(Category, related_name='articles')
""",
    )

    document.add_heading(
        "Paso 4: plantilla base con bloques reutilizables",
        level=2,
    )
    document.add_paragraph(
        "base.html define la estructura común, carga la hoja de estilos con "
        "{% load static %} y expone los bloques title, content y sidebar."
    )
    add_code(
        document,
        """
{% load static %}
<!DOCTYPE html>
<html lang="es">
<head>
    <title>{% block title %}Tecsup News Portal{% endblock %}</title>
    <link rel="stylesheet" href="{% static 'css/styles.css' %}">
</head>
<body>
    <header class="navbar">
        <a href="{% url 'news:home' %}" class="logo">Tecsup<span>Digital</span></a>
        <nav>
            <a href="{% url 'news:home' %}">Inicio</a>
            {% for category in categories %}
                <a href="{% url 'news:category_articles' category.slug %}">
                    {{ category.name }}
                </a>
            {% endfor %}
        </nav>
    </header>
    <main class="container layout">
        <section class="main-content">
            {% block content %}{% endblock %}
        </section>
        <aside class="sidebar">{% block sidebar %}...{% endblock %}</aside>
    </main>
    <footer>...</footer>
</body>
</html>
""",
    )

    document.add_heading(
        "Paso 5: fragmento reutilizable _article_card.html",
        level=2,
    )
    document.add_paragraph(
        "La tarjeta encapsula la imagen, la fecha, el título, el resumen, el "
        "autor y las categorías. Se incluye con {% include %} desde la portada "
        "y desde el listado por categoría."
    )
    add_code(
        document,
        """
<article class="article-card">
    {% if article.featured_image %}
        <img src="{{ article.featured_image.url }}" alt="{{ article.title }}">
    {% endif %}
    <div class="card-body">
        <span>{{ article.published_date|date:"d M Y" }}</span>
        <h3>
            <a href="{% url 'news:article_detail' article.slug %}">
                {{ article.title }}
            </a>
        </h3>
        <p>{{ article.summary|truncatewords:20 }}</p>
        <div>Por: {{ article.author }}</div>
        {% for category in article.categories.all %}
            <a href="{% url 'news:category_articles' category.slug %}">
                {{ category.name }}
            </a>
        {% endfor %}
    </div>
</article>
""",
    )

    document.add_heading(
        "Pasos 6, 7 y 8: portada, detalle y categoría",
        level=2,
    )
    add_code(
        document,
        """
{# home.html #}
{% extends 'base.html' %}
{% block title %}Portada - Portal de Noticias{% endblock %}
{% block content %}
    <h1>Últimas Noticias</h1>
    <div class="articles-grid">
        {% for article in articles %}
            {% include 'news/_article_card.html' with article=article %}
        {% empty %}
            <p>No hay noticias publicadas en este momento.</p>
        {% endfor %}
    </div>
{% endblock %}
""",
    )
    add_figure(
        document,
        "08_portada_sin_resultados.png",
        "Figura 2. Estado vacío renderizado con la etiqueta {% empty %}.",
    )
    document.add_paragraph(
        "El detalle hereda de base.html y muestra la imagen destacada, la "
        "fecha, el autor, el contenido y las categorías. El listado por "
        "categoría reutiliza el mismo fragmento de tarjeta."
    )
    add_figure(
        document,
        "02_detalle_noticia.png",
        "Figura 3. Detalle de noticia con imagen, autor y categorías.",
    )
    add_figure(
        document,
        "03_listado_categoria.png",
        "Figura 4. Listado por categoría con la tarjeta reutilizada.",
    )

    document.add_heading(
        "Pasos 9 y 10: rutas nombradas, vistas y estilos",
        level=2,
    )
    document.add_paragraph(
        'Las rutas usan app_name = "news" y nombres explícitos. Las plantillas '
        "generan todos los enlaces con {% url %}; no hay direcciones escritas "
        "a mano. La hoja de estilos se carga con {% static %}."
    )
    add_code(
        document,
        """
app_name = 'news'

urlpatterns = [
    path('', views.home_view, name='home'),
    path(
        'article/<slug:slug>/',
        views.article_detail_view,
        name='article_detail',
    ),
    path(
        'category/<slug:slug>/',
        views.category_list_view,
        name='category_articles',
    ),
]
""",
    )
    add_table(
        document,
        ["Recurso", "HTTP", "Content-Type"],
        [
            ("/", "200", "text/html; charset=utf-8"),
            ("/article/<slug>/", "200", "text/html; charset=utf-8"),
            ("/category/<slug>/", "200", "text/html; charset=utf-8"),
            ("/static/css/styles.css", "200", "text/css"),
            ("/media/articles/<archivo>.jpg", "200", "image/jpeg"),
        ],
        widths=[7.5, 2, 6],
    )

    document.add_heading(
        "Paso 11: administración personalizada y carga de datos",
        level=2,
    )
    document.add_paragraph(
        "Se configuraron list_display, list_filter, search_fields, "
        "prepopulated_fields, filter_horizontal y date_hierarchy. Se cargaron "
        "seis noticias en tres categorías (Tecnología, Educación e Innovación) "
        "con tres autores. Las imágenes se generaron con Pillow mediante el "
        "comando seed_demo."
    )
    add_code(
        document,
        """
@admin.register(Article)
class ArticleAdmin(admin.ModelAdmin):
    list_display = (
        'title', 'author', 'published_date', 'is_published', 'category_list',
    )
    list_filter = ('is_published', 'published_date', 'categories', 'author')
    search_fields = ('title', 'summary', 'content', 'author__last_name')
    prepopulated_fields = {'slug': ('title',)}
    filter_horizontal = ('categories',)
    date_hierarchy = 'published_date'
""",
    )
    add_figure(
        document,
        "05_admin_articulos.png",
        "Figura 5. Admin de artículos con columnas, filtros y búsqueda.",
    )
    add_figure(
        document,
        "06_admin_categorias.png",
        "Figura 6. Admin de categorías con conteo de artículos.",
    )
    add_figure(
        document,
        "07_admin_autores.png",
        "Figura 7. Admin de autores con conteo de artículos.",
    )

    document.add_heading(
        "Paso 12: escapado automático de HTML",
        level=2,
    )
    document.add_paragraph(
        "Se guardó en el cuerpo de una noticia el texto <b>Texto en negrita</b> "
        'y <script>alert("XSS")</script>. La plantilla lo renderiza con '
        "{{ article.content }} sin el filtro |safe."
    )
    document.add_paragraph(
        "Django aplica auto-escaping y convierte < en &lt;, > en &gt; y las "
        "comillas en &quot;. El navegador muestra el código como texto visible "
        "y no ejecuta el script. Esto evita ataques de Cross-Site Scripting "
        "(XSS). Para mostrar HTML confiable se requeriría sanearlo y usar "
        "explícitamente |safe o {% autoescape off %}, opción no recomendada "
        "para contenido de usuarios."
    )
    add_figure(
        document,
        "04_escapado_html.png",
        "Figura 8. Etiquetas HTML y script mostrados literalmente por Django.",
    )

    document.add_heading(
        "Paso 13: repositorio y entregable",
        level=2,
    )
    document.add_paragraph(
        "El entregable incluye el código fuente, este informe, ocho capturas "
        "de pantalla, diez pruebas automatizadas y el comando seed_demo para "
        "reproducir los datos e imágenes. La subida al repositorio del equipo "
        "y al campus virtual se realiza con las credenciales institucionales."
    )

    document.add_heading("4. Matriz de casos de prueba", level=1)
    add_table(
        document,
        ["ID", "Caso de prueba", "Entrada / acción", "Resultado esperado", "Estado"],
        [
            (
                "CP-01",
                "Herencia de plantillas",
                "Navegar entre portada y detalle",
                "La estructura común se mantiene sin duplicar HTML",
                "PASÓ",
            ),
            (
                "CP-02",
                "Reutilización de fragmentos",
                "Cargar portada y categoría",
                "Ambas renderizan _article_card.html",
                "PASÓ",
            ),
            (
                "CP-03",
                "Gestión desde el admin",
                "Modificar el título de una noticia",
                "El cambio se refleja de inmediato en el portal",
                "PASÓ",
            ),
            (
                "CP-04",
                "Escapado HTML",
                "Guardar <b> y <script> en el contenido",
                "Se muestran como texto escapado; no se ejecuta XSS",
                "PASÓ",
            ),
            (
                "CP-05",
                "Estado vacío",
                "Despublicar todas las noticias",
                "Se muestra el bloque {% empty %}",
                "PASÓ",
            ),
            (
                "CP-06",
                "Filtros de plantilla",
                "Consultar la portada",
                "Fecha con date y resumen con truncatewords",
                "PASÓ",
            ),
            (
                "CP-07",
                "Rutas nombradas",
                "Revertir las tres rutas",
                "URLs correctas",
                "PASÓ",
            ),
            (
                "CP-08",
                "Archivos estáticos",
                "Buscar css/styles.css",
                "Archivo encontrado y servido con HTTP 200",
                "PASÓ",
            ),
            (
                "CP-09",
                "Noticia no publicada",
                "Abrir su detalle",
                "Respuesta HTTP 404",
                "PASÓ",
            ),
            (
                "CP-10",
                "Relaciones en el detalle",
                "Abrir una noticia",
                "Se muestran imagen, autor y categorías",
                "PASÓ",
            ),
        ],
        widths=[1.5, 4.1, 4.1, 5.2, 1.6],
    )
    add_code(
        document,
        """
Found 10 test(s).
..........
----------------------------------------------------------------------
Ran 10 tests

OK
""",
    )

    document.add_heading(
        "5. Autoevaluación según la rúbrica del curso (20 puntos)",
        level=1,
    )
    add_table(
        document,
        ["Criterio", "Implementación lograda", "Nivel", "Puntaje"],
        [
            (
                "1. Herencia y fragmentos reutilizables",
                "base.html con bloques y _article_card.html reutilizado sin "
                "duplicación",
                "Excelente",
                "5 / 5",
            ),
            (
                "2. Variables, etiquetas y filtros",
                "{% for %}, {% empty %}, date y truncatewords aplicados en plantillas",
                "Excelente",
                "5 / 5",
            ),
            (
                "3. Gestión de contenido desde el admin",
                "CRUD de noticias, autores y categorías con panel personalizado",
                "Excelente",
                "5 / 5",
            ),
            (
                "4. Entregable y repositorio",
                "Estructura documentada, código en inglés, capturas y pruebas",
                "Excelente",
                "5 / 5",
            ),
        ],
        widths=[5.2, 7.2, 2.4, 1.8],
    )

    document.add_heading("6. Conclusiones técnicas", level=1)
    for text in (
        'La herencia con {% extends "base.html" %} mantiene un diseño '
        "homogéneo y evita duplicar el marcado; un cambio en la base se "
        "propaga a todas las páginas.",
        "Los fragmentos con {% include %} permiten reutilizar la tarjeta de "
        "noticia en la portada y en los listados por categoría, reduciendo "
        "inconsistencias.",
        "La separación entre modelos, vistas y plantillas mantiene la lógica "
        "de negocio fuera de la capa de presentación.",
        "El auto-escapado de Django neutraliza HTML y scripts no confiables y "
        "constituye una defensa esencial contra XSS.",
        "El comando seed_demo hace reproducible la demostración: genera las "
        "imágenes con Pillow y carga categorías, autores y noticias.",
    ):
        document.add_paragraph(text, style="List Number")

    document.add_heading("Anexo: comandos de ejecución", level=1)
    add_code(
        document,
        """
python -m pip install -r requirements.txt
python manage.py migrate
python manage.py seed_demo
python manage.py createsuperuser
python manage.py runserver
python manage.py test
""",
    )

    document.save(OUTPUT)
    print(f"Report saved at {OUTPUT}")


if __name__ == "__main__":
    build()
